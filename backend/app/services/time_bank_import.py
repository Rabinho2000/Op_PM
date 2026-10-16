"""Bounded, non-evaluating XLSX parser for the Banco de Horas import.

Only cached values and the small set of formula shapes observed in the source
workbook are accepted.  The workbook itself never enters the database: the
staging JSON contains deterministic row data and provenance only.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl.utils.exceptions import InvalidFileException

from app.services.time_bank import calculate_work_seconds

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_MEMBERS = 512
PARSER_VERSION = "xlsx-1"
MONTH_SHEETS = (
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
)
_SECRET_NOTE_PATTERN = re.compile(
    r"(?i)(?:https?://|data:|bearer\s+|(?:api[_-]?key|authorization|password|secret|token)\s*[:=]|-----begin)"
)


@dataclass(frozen=True)
class ImportIssueData:
    code: str
    message: str
    blocking: bool
    sheet: str | None = None
    row: int | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "code": self.code,
            "message": self.message,
            "blocking": self.blocking,
            "sheet": self.sheet,
            "row": self.row,
        }


class TimeBankImportError(ValueError):
    def __init__(self, message: str, *, code: str = "invalid_workbook") -> None:
        super().__init__(message)
        self.code = code


def _safe_text(value: object, *, limit: int = 512) -> str:
    if value is None:
        return ""
    text = str(value).replace("\x00", "").strip()
    text = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    return text[:limit]


def _safe_note(value: object) -> tuple[str | None, bool]:
    """Keep human notes while removing URL/secret-shaped payloads."""
    text = _safe_text(value, limit=2000)
    if not text:
        return None, False
    if not _SECRET_NOTE_PATTERN.search(text):
        return text, False
    sanitised = re.sub(r"(?i)https?://[^\s]+", "[URL omitido]", text)
    sanitised = re.sub(r"(?i)data:[^\s]+", "[payload omitido]", sanitised)
    sanitised = re.sub(r"(?i)bearer\s+[^\s]+", "Bearer [omitido]", sanitised)
    sanitised = re.sub(
        r"(?i)((?:api[_-]?key|authorization|password|secret|token)\s*[:=])\s*[^\s]+",
        r"\1[omitido]",
        sanitised,
    )
    sanitised = re.sub(r"(?is)-----begin.*", "[conteúdo sensível omitido]", sanitised)
    return sanitised, True


def _row_notes(formula_row: tuple[Any, ...], cached_row: tuple[Any, ...]) -> tuple[list[dict[str, Any]], bool]:
    notes: list[dict[str, Any]] = []
    redacted = False
    seen: set[tuple[int, str]] = set()

    def add_note(column: int, value: object, kind: str) -> None:
        nonlocal redacted
        text, was_redacted = _safe_note(value)
        if text is None:
            return
        key = (column, text)
        if key in seen:
            return
        seen.add(key)
        redacted = redacted or was_redacted
        notes.append({"column": column, "kind": kind, "text": text})

    # Values outside the canonical A:G table are imported as notes. Comments
    # attached to G are also source notes in the workbook and must not vanish.
    for column, cell in enumerate(cached_row[7:], start=8):
        if cell.value is not None:
            add_note(column, cell.value, "cell")
    for column, cell in enumerate(cached_row, start=1):
        comment = getattr(cell, "comment", None)
        if comment is not None:
            add_note(column, comment.text, "comment")
    # Some openpyxl versions expose comments only on the formula workbook.
    if not any(getattr(cell, "comment", None) is not None for cell in cached_row):
        for column, cell in enumerate(formula_row, start=1):
            comment = getattr(cell, "comment", None)
            if comment is not None:
                add_note(column, comment.text, "comment")
    return notes, redacted


def _zip_safety_check(content: bytes) -> None:
    if len(content) > MAX_UPLOAD_BYTES:
        raise TimeBankImportError("O ficheiro excede o limite de 5 MiB.", code="file_too_large")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ZIP_MEMBERS:
                raise TimeBankImportError("O XLSX contém demasiados componentes.", code="zip_bomb")
            total = 0
            seen: set[str] = set()
            for info in infos:
                name = info.filename
                if name in seen or name.startswith("/") or ".." in name.split("/"):
                    raise TimeBankImportError("Estrutura ZIP inválida.", code="invalid_zip")
                seen.add(name)
                if info.file_size > MAX_UNCOMPRESSED_BYTES or info.file_size < 0:
                    raise TimeBankImportError("Um componente XML excede o limite permitido.", code="zip_bomb")
                total += info.file_size
                if total > MAX_UNCOMPRESSED_BYTES:
                    raise TimeBankImportError("O conteúdo XML descomprimido excede o limite permitido.", code="zip_bomb")
            for info in infos:
                if not info.filename.endswith(".rels"):
                    continue
                raw = archive.read(info)
                if b"TargetMode=\"External\"" in raw or b"TargetMode='External'" in raw:
                    raise TimeBankImportError("Relações externas não são aceites.", code="external_relation")
    except zipfile.BadZipFile as exc:
        raise TimeBankImportError("O ficheiro não é um XLSX válido.", code="invalid_xlsx") from exc


def _normalise_header(value: object) -> str:
    return re.sub(r"\s+", " ", _safe_text(value).casefold())


def _formula_is_known(formula: object) -> bool:
    if not isinstance(formula, str) or not formula.startswith("="):
        return True
    upper = formula.upper()
    if "[" in upper and "TABLE_" not in upper:
        return False
    if any(token in upper for token in ("EXTERNAL", "WEBSERVICE", "INDIRECT", "OFFSET", "HYPERLINK", "_XLFN")):
        return False
    functions = set(re.findall(r"([A-Z][A-Z0-9_.]*)\s*\(", upper))
    return functions.issubset({"IF", "WEEKDAY", "SUM", "SUMIF"}) and re.fullmatch(
        r"[=\w.$!:#,;\-+*/<>=() \"'\[\]ÁÉÍÓÚÀÃÕÂÊÔÇáéíóúàãõâêôç]+", upper
    ) is not None


def _as_date(value: object) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return None


def _hours_to_minutes(value: object, *, field: str) -> tuple[int | None, str | None]:
    if value is None or value == "":
        return None, None
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, AttributeError):
        return None, f"{field}_invalid"
    if number < 0 or number > 24 or (number == 24 and field != "end_minute"):
        return None, f"{field}_out_of_range"
    minutes = number * 60
    if minutes != minutes.to_integral_value():
        return None, f"{field}_not_whole_minute"
    return int(minutes), None


def _hours_to_seconds(value: object) -> int | None:
    if value is None or value == "":
        return None
    try:
        seconds = Decimal(str(value).strip()) * 3600
    except (InvalidOperation, AttributeError):
        return None
    if seconds != seconds.to_integral_value():
        return None
    return int(seconds)


def _lunch_value(value: object) -> bool | None:
    text = _safe_text(value).casefold()
    if text in {"sim", "yes", "true", "1"}:
        return True
    if text in {"não", "nao", "no", "false", "0"}:
        return False
    return None


def _issue(issues: list[ImportIssueData], code: str, message: str, *, blocking: bool, sheet: str, row: int) -> None:
    issues.append(ImportIssueData(code, message, blocking, sheet, row))


def _find_daily_header(ws: Any) -> int | None:
    wanted = {"colaborador", "weekday", "data", "hora de entrada", "hora de almoço", "hora de saída", "balanço"}
    for row_number, row in enumerate(ws.iter_rows(min_row=1, max_row=40, values_only=True), start=1):
        values = {_normalise_header(value) for value in row[:8] if value is not None}
        if wanted.issubset(values):
            return row_number
    return None


def _source_key(file_hash: str, kind: str, sheet: str, row: int) -> str:
    return f"{file_hash}:{kind}:{sheet}:{row}"


def _parse_daily_sheet(
    formula_ws: Any,
    cached_ws: Any,
    *,
    file_hash: str,
    issues: list[ImportIssueData],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], int]:
    header = _find_daily_header(cached_ws)
    if header is None:
        _issue(issues, "missing_daily_header", "Folha mensal sem cabeçalho reconhecido.", blocking=True, sheet=cached_ws.title, row=1)
        return [], [], 0
    entries: list[dict[str, Any]] = []
    quarantined: list[dict[str, Any]] = []
    source_rows = 0
    blank_streak = 0
    formula_rows = formula_ws.iter_rows(min_row=header + 1)
    cached_rows = cached_ws.iter_rows(min_row=header + 1)
    for offset, (frow, crow) in enumerate(zip(formula_rows, cached_rows), start=header + 1):
        cached_values = [cell.value for cell in crow]
        formula_values = [cell.value for cell in frow]
        if not any(value is not None for value in cached_values[:10]):
            blank_streak += 1
            if blank_streak >= 25:
                break
            continue
        label = _safe_text(cached_values[0] if cached_values else None)
        # The monthly template contains thousands of formula-only rows. They
        # are not source rows; only a labelled line is an input identity that
        # must be reconciled.
        if not label:
            continue
        source_rows += 1
        blank_streak = 0
        row_date = _as_date(cached_values[2] if len(cached_values) > 2 else None)
        flags: dict[str, Any] = {}
        row_notes, notes_redacted = _row_notes(frow, crow)
        if notes_redacted:
            flags["note_redacted"] = True
            _issue(issues, "note_redacted", "Uma nota com formato sensível/URL foi omitida da staging.", blocking=False, sheet=cached_ws.title, row=offset)
        notes = " | ".join(note["text"] for note in row_notes) or None

        if row_date is None:
            balance_cached = _hours_to_seconds(cached_values[6] if len(cached_values) > 6 else None)
            flags["undated"] = True
            if balance_cached is None:
                flags["value_not_recoverable"] = True
                _issue(issues, "undated_value_not_recoverable", "Linha identificada sem data e sem balanço numérico; permanece em quarentena.", blocking=True, sheet=cached_ws.title, row=offset)
            elif balance_cached == 0:
                _issue(issues, "undated_zero_quarantined", "Linha identificada sem data, com valor zero; preservada em quarentena sem criar data.", blocking=False, sheet=cached_ws.title, row=offset)
            else:
                _issue(issues, "undated_nonzero", "Linha identificada sem data e com valor não zero; resolução manual obrigatória.", blocking=True, sheet=cached_ws.title, row=offset)
            quarantined.append(
                {
                    "source_key": _source_key(file_hash, "quarantine", cached_ws.title, offset),
                    "source_label": label,
                    "work_date": None,
                    "balance_seconds": balance_cached,
                    "reason": notes,
                    "needs_review": True,
                    "review_flags": flags,
                    "provenance": {
                        "source_hash": file_hash,
                        "sheet": cached_ws.title,
                        "row": offset,
                        "notes": row_notes,
                    },
                }
            )
            continue
        weekday_cached = cached_values[1] if len(cached_values) > 1 else None
        weekday_formula = formula_values[1] if len(formula_values) > 1 else None
        if not _formula_is_known(weekday_formula):
            flags["unknown_weekday_formula"] = True
            _issue(issues, "unknown_formula", "Fórmula Weekday não reconhecida; o valor em cache foi preservado.", blocking=weekday_cached is None, sheet=cached_ws.title, row=offset)
        expected_weekday = ((row_date.weekday() + 1) % 7) + 1  # Excel WEEKDAY(date) default: Sunday=1
        if isinstance(weekday_cached, (int, float)) and int(weekday_cached) != expected_weekday:
            flags["weekday_divergent"] = {"source": int(weekday_cached), "calculated": expected_weekday}
            _issue(issues, "weekday_divergent", "Weekday divergente da data; o valor histórico não foi corrigido.", blocking=False, sheet=cached_ws.title, row=offset)

        start, start_issue = _hours_to_minutes(cached_values[3] if len(cached_values) > 3 else None, field="start_minute")
        end, end_issue = _hours_to_minutes(cached_values[5] if len(cached_values) > 5 else None, field="end_minute")
        if start_issue:
            flags[start_issue] = True
            _issue(issues, start_issue, "Hora de entrada inválida; valor histórico não recuperável.", blocking=True, sheet=cached_ws.title, row=offset)
        if end_issue:
            flags[end_issue] = True
            _issue(issues, end_issue, "Hora de saída inválida; valor histórico não recuperável.", blocking=True, sheet=cached_ws.title, row=offset)
        lunch = _lunch_value(cached_values[4] if len(cached_values) > 4 else None)
        if cached_values[4] is not None and lunch is None:
            flags["lunch_unknown"] = True
            _issue(issues, "lunch_unknown", "Valor de almoço não reconhecido; valor histórico preservado.", blocking=False, sheet=cached_ws.title, row=offset)

        balance_formula = formula_values[6] if len(formula_values) > 6 else None
        balance_cached = _hours_to_seconds(cached_values[6] if len(cached_values) > 6 else None)
        formula_known = _formula_is_known(balance_formula)
        if not formula_known:
            flags["unknown_balance_formula"] = True
            _issue(issues, "unknown_formula", "Fórmula de balanço não reconhecida; apenas o cache numérico pode ser aplicado.", blocking=balance_cached is None, sheet=cached_ws.title, row=offset)
        mode = "imported"
        calculated: int | None = None
        if start is not None and end is not None and lunch is not None and not flags.get("weekday_divergent"):
            calc_mode = "saturday" if row_date.weekday() == 5 else "standard" if row_date.weekday() < 5 else "manual"
            if calc_mode != "manual":
                try:
                    calculated = calculate_work_seconds(
                        work_date=row_date, start_minute=start, end_minute=end,
                        lunch_taken=lunch, calculation_mode=calc_mode,
                    )
                except ValueError:
                    flags["calculation_not_recoverable"] = True
            else:
                flags["sunday_manual"] = True
        elif start is None or end is None or lunch is None:
            flags["missing_hours"] = True
            _issue(issues, "missing_hours", "Horário incompleto; o lançamento só pode ser aplicado com valor histórico em cache.", blocking=balance_cached is None, sheet=cached_ws.title, row=offset)
        if balance_cached is None and calculated is not None and formula_known:
            balance_cached = calculated
            flags["formula_cache_missing_recomputed"] = True
            _issue(issues, "formula_cache_missing", "Cache ausente; valor standard foi recalculado e marcado para revisão.", blocking=False, sheet=cached_ws.title, row=offset)
        elif balance_cached is not None and calculated is not None and balance_cached != calculated:
            flags["formula_value_mismatch"] = {"calculated_seconds": calculated}
            _issue(issues, "formula_value_mismatch", "Valor histórico difere do cálculo atual; ambos foram preservados.", blocking=False, sheet=cached_ws.title, row=offset)
        if balance_cached is None:
            flags["value_not_recoverable"] = True
            _issue(issues, "value_not_recoverable", "Balanço sem cache numérico e sem cálculo seguro; aplicação bloqueada.", blocking=True, sheet=cached_ws.title, row=offset)

        entries.append(
            {
                "source_key": _source_key(file_hash, "entry", cached_ws.title, offset),
                "source_label": label,
                "work_date": row_date.isoformat(),
                "start_minute": start,
                "end_minute": end,
                "lunch_taken": lunch,
                "balance_seconds": balance_cached,
                "calculated_seconds": calculated,
                "calculation_mode": mode,
                "reason": notes,
                "needs_review": bool({key: value for key, value in flags.items() if key != "note_redacted"}),
                "review_flags": flags,
                "provenance": {
                    "source_hash": file_hash,
                    "sheet": cached_ws.title,
                    "row": offset,
                    "formula": "known" if formula_known else "unknown",
                    "notes": row_notes,
                },
            }
        )
    return entries, quarantined, source_rows


def _column_name(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _summary_layout(cached_ws: Any, source_year: int) -> tuple[list[tuple[int, str]], set[str]]:
    row1 = list(cached_ws[1])
    row2 = list(cached_ws[2])
    year_columns = {
        index + 1: cell.value
        for index, cell in enumerate(row1)
        if isinstance(cell.value, int) and 1900 <= cell.value <= 2200
    }
    current_column = next((column for column, year in year_columns.items() if year == source_year), None)
    previous_column = next((column for column, year in year_columns.items() if year == source_year - 1), None)
    if current_column is None:
        return [], set()

    current_labels: list[tuple[int, str]] = []
    current_limit = previous_column or (min(year_columns) if year_columns else len(row2) + 1)
    for column in range(current_column + 1, current_limit):
        value = _safe_text(row2[column - 1].value if column <= len(row2) else None)
        if not value:
            break
        if value.casefold() not in {"mês", "mes", "total"}:
            current_labels.append((column, value))

    allowed_previous: set[str] = set()
    if previous_column is not None:
        next_year_column = min((column for column in year_columns if column > previous_column), default=len(row2) + 1)
        for column in range(previous_column + 1, next_year_column):
            value = _safe_text(row2[column - 1].value if column <= len(row2) else None)
            if value and value.casefold() not in {"mês", "mes", "total"}:
                allowed_previous.add(_column_name(column))
    return current_labels, allowed_previous


def _parse_openings(
    formula_ws: Any,
    cached_ws: Any,
    *,
    file_hash: str,
    source_year: int | None,
    issues: list[ImportIssueData],
) -> list[dict[str, Any]]:
    if source_year is None:
        return []
    current_labels, allowed_previous = _summary_layout(cached_ws, source_year)
    formulas = [cell.value for cell in formula_ws[15]]
    openings: list[dict[str, Any]] = []
    # Only the current-year total formula may reference the immediately
    # previous-year block. Older blocks and plain numeric historical totals are
    # never added to the opening balance.
    for column, label in current_labels:
        formula = formulas[column - 1] if column <= len(formulas) else None
        current_column = _column_name(column)
        if not isinstance(formula, str) or not formula.startswith("=") or not _formula_is_known(formula):
            continue
        normalised = formula.upper().replace("$", "").replace(" ", "")
        prefix = f"=SUM({current_column}3:{current_column}14)"
        if not normalised.startswith(prefix):
            continue
        tail = normalised[len(prefix):]
        if not re.fullmatch(r"(?:\+[A-Z]{1,3}15)+", tail):
            continue
        references = re.findall(r"\+([A-Z]{1,3})15", tail)
        for col in references:
            if col not in allowed_previous:
                continue
            ref_value = cached_ws[f"{col}15"].value
            delta = _hours_to_seconds(ref_value)
            if delta is None:
                _issue(issues, "opening_value_not_recoverable", "Saldo transitado referenciado sem valor numérico recuperável.", blocking=True, sheet=cached_ws.title, row=15)
                continue
            if delta <= 0:
                continue
            openings.append(
                {
                    "source_key": f"{file_hash}:opening:{col}15",
                    "source_label": label,
                    "effective_date": f"{source_year:04d}-01-01",
                    "delta_seconds": delta,
                    "kind": "opening",
                    "reason": "Saldo transitado referenciado pelo resumo anual.",
                    "provenance": {
                        "source_hash": file_hash,
                        "sheet": cached_ws.title,
                        "row": 15,
                        "reference": f"{col}15",
                    },
                }
            )
    return openings


def parse_time_bank_xlsx(content: bytes) -> dict[str, Any]:
    """Parse an XLSX into sanitized JSON-like staging data."""
    _zip_safety_check(content)
    file_hash = hashlib.sha256(content).hexdigest()
    try:
        import openpyxl

        formula_wb = openpyxl.load_workbook(io.BytesIO(content), read_only=False, data_only=False, keep_links=False)
        cached_wb = openpyxl.load_workbook(io.BytesIO(content), read_only=False, data_only=True, keep_links=False)
    except (
        ET.ParseError,
        InvalidFileException,
        OSError,
        ValueError,
        KeyError,
        IndexError,
        zipfile.BadZipFile,
        zipfile.LargeZipFile,
    ) as exc:
        raise TimeBankImportError("Não foi possível ler a estrutura XLSX.", code="invalid_xlsx") from exc

    issues: list[ImportIssueData] = []
    entries: list[dict[str, Any]] = []
    quarantined: list[dict[str, Any]] = []
    source_rows = 0
    for sheet in MONTH_SHEETS:
        if sheet not in formula_wb.sheetnames or sheet not in cached_wb.sheetnames:
            _issue(issues, "missing_month_sheet", "Falta uma folha mensal esperada.", blocking=False, sheet=sheet, row=1)
            continue
        parsed_entries, parsed_quarantine, parsed_source_rows = _parse_daily_sheet(
            formula_wb[sheet], cached_wb[sheet], file_hash=file_hash, issues=issues
        )
        entries.extend(parsed_entries)
        quarantined.extend(parsed_quarantine)
        source_rows += parsed_source_rows

    source_year: int | None = None
    openings: list[dict[str, Any]] = []
    summary_reconciliation: list[dict[str, Any]] = []
    if "Balanço final" in formula_wb.sheetnames and "Balanço final" in cached_wb.sheetnames:
        formula_summary = formula_wb["Balanço final"]
        cached_summary = cached_wb["Balanço final"]
        first = cached_summary[1][0].value
        if isinstance(first, int) and 1900 <= first <= 2200:
            source_year = first
        openings = _parse_openings(formula_summary, cached_summary, file_hash=file_hash, source_year=source_year, issues=issues)
        if source_year is not None:
            current_labels, _ = _summary_layout(cached_summary, source_year)
            daily_totals: defaultdict[str, int] = defaultdict(int)
            for row in entries:
                if isinstance(row.get("balance_seconds"), int):
                    daily_totals[row["source_label"]] += row["balance_seconds"]
            opening_totals: defaultdict[str, int] = defaultdict(int)
            for row in openings:
                opening_totals[row["source_label"]] += row["delta_seconds"]
            for column, label in current_labels:
                summary_value = _hours_to_seconds(cached_summary.cell(15, column).value)
                daily_value = daily_totals.get(label, 0)
                opening_value = opening_totals.get(label, 0)
                ledger_value = daily_value + opening_value
                flags: list[str] = []
                if label not in daily_totals and label not in opening_totals:
                    flags.append("summary_without_daily_rows")
                    _issue(issues, "summary_person_without_daily_rows", "Pessoa no cabeçalho do ano corrente sem linhas diárias; não foi criada uma pessoa nova.", blocking=False, sheet=cached_summary.title, row=2)
                if summary_value is None:
                    flags.append("summary_value_not_recoverable")
                    _issue(issues, "summary_value_not_recoverable", "Total do resumo anual sem valor numérico recuperável.", blocking=True, sheet=cached_summary.title, row=15)
                    difference = None
                else:
                    difference = summary_value - ledger_value
                    if difference != 0:
                        flags.append("summary_total_mismatch")
                        _issue(issues, "summary_total_mismatch", "Total do resumo difere das linhas diárias mais o saldo transitado; ambos os valores foram preservados, sem autocorreção.", blocking=False, sheet=cached_summary.title, row=15)
                summary_reconciliation.append(
                    {
                        "source_label": label,
                        "summary_seconds": summary_value,
                        "daily_seconds": daily_value,
                        "opening_seconds": opening_value,
                        "ledger_seconds": ledger_value,
                        "difference_seconds": difference,
                        "flags": flags,
                    }
                )
    else:
        _issue(issues, "missing_summary_sheet", "Falta a folha de balanço final.", blocking=True, sheet="Balanço final", row=1)

    people: list[str] = []
    for row in entries + openings + quarantined:
        label = row["source_label"]
        if label not in people:
            people.append(label)

    issue_dicts = [item.as_dict() for item in issues]
    return {
        "parser_version": PARSER_VERSION,
        "source_hash": file_hash,
        "source_year": source_year,
        "entries": entries,
        "openings": openings,
        "quarantined": quarantined,
        "people": people,
        "summary_reconciliation": summary_reconciliation,
        "issues": issue_dicts,
        "counts": {
            "entries": len(entries),
            "openings": len(openings),
            "needs_review": sum(1 for row in entries if row["needs_review"]),
            "source_rows": source_rows,
            "quarantined": len(quarantined),
        },
    }
