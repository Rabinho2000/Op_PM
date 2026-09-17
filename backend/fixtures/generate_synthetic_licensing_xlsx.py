#!/usr/bin/env python3
"""Gera `synthetic_licensing.xlsx` — fixture sintética para testes do
importador de licenciamento (app/services/imports_licensing.py). Nunca
contém dados reais. Corre-se uma vez; o ficheiro gerado é commitado.

    python fixtures/generate_synthetic_licensing_xlsx.py

Inclui de propósito:
- `Project_number` repetido (1001 em duas linhas) com `Internal_reference`
  diferente — prova que a chave de correspondência é sempre
  `Internal_reference`, nunca `Project_number` isolado.
- Colunas de credenciais (PIN, PUK, Portal_password, Portal_login) na
  folha "Dados gerais" — o importador tem de as ignorar sempre.
- Uma linha em "Venda do excedente" sem projeto correspondente — prova a
  fila de conflitos/não-ligados.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook

OUTPUT_PATH = Path(__file__).resolve().parent / "synthetic_licensing.xlsx"


def build() -> Workbook:
    wb = Workbook()

    sheet1 = wb.active
    sheet1.title = "Sheet1"
    sheet1.append(
        [
            "Project_number",
            "Project_name",
            "Client_name",
            "NIF",
            "Type",
            "Internal_reference",
            "Date",
            "Adress",
            "Location",
            "Panel_count",
            "Panel_power_wp",
            "Power_kwp",
            "Installation_type",
            "Injection",
            "Batteries",
            "Chargers",
            "Vendedor",
            "OM",
            "Estado",
            "Contacts",
            "Notes",
            "Contracted_month_year",
        ]
    )
    sheet1.append(
        [
            1001,
            "Instalação Sintética Excel A",
            "Cliente Sintético Excel A",
            "500111222",
            "Residencial",
            "REF-EXCEL-001",
            dt.date(2025, 3, 1),
            "Rua Sintética Excel, 1",
            "Porto",
            20,
            450,
            9.0,
            "Autoconsumo",
            "Injeção parcial",
            "Sem bateria",
            "Nenhum",
            "Vendedor Sintético A",
            "Sem contrato de O&M",
            "Concluído",
            "Contacto Sintético A",
            "Nota sintética A.",
            "03/2025",
        ]
    )
    sheet1.append(
        [
            # Mesmo Project_number que a linha anterior, de propósito —
            # prova que a chave de correspondência é Internal_reference.
            1001,
            "Instalação Sintética Excel B",
            "Cliente Sintético Excel B",
            "500333444",
            "Comercial",
            "REF-EXCEL-002",
            dt.date(2025, 6, 15),
            "Rua Sintética Excel, 2",
            "Braga",
            40,
            465,
            18.6,
            "Autoconsumo com armazenamento",
            "Injeção total",
            "1x bateria sintética",
            "1x carregador sintético",
            "Vendedor Sintético B",
            "Contrato de O&M trienal sintético",
            "Em curso",
            "Contacto Sintético B",
            "Nota sintética B.",
            "06/2025",
        ]
    )

    dados_gerais = wb.create_sheet("Dados gerais")
    dados_gerais.append(
        [
            "Internal_reference",
            "M2M_number",
            "Operator",
            "UPAC_number",
            "Registration",
            "Cadastro",
            "Licensing_status",
            "Inspecting_entity",
            "Inspection_date",
            "Certificate_date",
            "Email",
            "Power_kwp",
            "Modules",
            "Inverters",
            "Installer",
            "Production_kwh",
            "Location",
            "Contract",
            "Contact",
            "Comments",
            # Colunas de credenciais — nunca devem ser importadas.
            "PIN",
            "PUK",
            "Portal_password",
            "Portal_login",
        ]
    )
    dados_gerais.append(
        [
            "REF-EXCEL-001",
            "912345000",
            "Operador Sintético X",
            "UPAC-EXCEL-001",
            "Registado",
            "CAD-001",
            "Registado",
            "Entidade Sintética de Inspeção",
            dt.date(2025, 4, 1),
            dt.date(2025, 4, 15),
            "cliente.excel.a@example.invalid",
            9.0,
            20,
            "1x inversor sintético 9kW",
            "Instalador Sintético A",
            14000,
            "Porto",
            "Contrato sintético A",
            "Contacto Sintético A",
            "Comentário sintético A.",
            "1234",
            "87654321",
            "segredo-sintetico-a",
            "utilizador-sintetico-a",
        ]
    )
    dados_gerais.append(
        [
            "REF-EXCEL-002",
            "913345111",
            "Operador Sintético Y",
            "UPAC-EXCEL-002",
            "Registado",
            "CAD-002",
            "Registado",
            "Entidade Sintética de Inspeção",
            dt.date(2025, 7, 1),
            dt.date(2025, 7, 20),
            "cliente.excel.b@example.invalid",
            18.6,
            40,
            "1x inversor trifásico sintético 19kW",
            "Instalador Sintético B",
            26000,
            "Braga",
            "Contrato sintético B",
            "Contacto Sintético B",
            "Comentário sintético B.",
            "5678",
            "12348765",
            "segredo-sintetico-b",
            "utilizador-sintetico-b",
        ]
    )

    surplus = wb.create_sheet("Venda do excedente")
    surplus.append(
        [
            "Internal_reference",
            "Commercializer",
            "Contract_type",
            "Status",
            "Sent_date",
            "Signed_date",
            "Duration",
            "Start_date",
            "End_date",
            "Notes",
        ]
    )
    surplus.append(
        [
            "REF-EXCEL-001",
            "Comercializador Sintético X",
            "Fixo",
            "Assinado",
            dt.date(2025, 4, 10),
            dt.date(2025, 4, 20),
            "12 meses",
            dt.date(2025, 5, 1),
            dt.date(2026, 5, 1),
            "Nota sintética de excedente A.",
        ]
    )
    surplus.append(
        [
            # Sem projeto correspondente de propósito — prova a fila de
            # "sem projeto identificado".
            "REF-EXCEL-999",
            "Comercializador Sintético Z",
            "Indexado",
            "Enviado",
            dt.date(2025, 8, 1),
            None,
            "24 meses",
            None,
            None,
            "Nota sintética de excedente sem projeto correspondente.",
        ]
    )

    return wb


if __name__ == "__main__":
    workbook = build()
    workbook.save(OUTPUT_PATH)
    print(f"Fixture gerada: {OUTPUT_PATH}")
