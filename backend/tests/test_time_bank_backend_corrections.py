from __future__ import annotations

import datetime as dt
import io
import json
import threading
import uuid
import zipfile

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import event
from sqlalchemy.orm import Session as OrmSession

from app.db import SessionLocal, engine
from app.main import app
from app.models.identity import User
from app.models.people import Person
from app.models.time_bank import (
    TimeBankAdjustment,
    TimeBankEntry,
    TimeBankImportBatch,
)
from app.security.permissions import AuthContext
from app.services.time_bank_import import MAX_UPLOAD_BYTES, PARSER_VERSION
from app.services.time_bank_ledger import (
    TimeBankError,
    apply_import,
    preview_import,
    resolve_import,
)


ADMIN = "admin.sintetico@example.invalid"
CHEFE = "chefe.sintetico@example.invalid"
UPLOAD_HEADERS = {"X-Dev-User-Email": ADMIN}


def _workbook_bytes(label: str = "Person Alpha", *, unsafe_formula: bool = False) -> bytes:
    workbook = openpyxl.Workbook()
    monthly = workbook.active
    monthly.title = "Janeiro"
    monthly.append(["Janeiro"])
    monthly.append(["Colaborador", "Weekday", "Data", "Hora de entrada", "Hora de almoço", "Hora de saída", "Balanço"])
    monthly.append([label, 2, dt.datetime(2026, 1, 5), 9, "sim", 18.5, 0])
    if unsafe_formula:
        monthly["G3"] = '=WEBSERVICE("https://example.invalid")'
    summary = workbook.create_sheet("Balanço final")
    summary.append([2026])
    summary.append(["Mês", label])
    for _ in range(12):
        summary.append(["Mês", 0])
    summary.append(["Total", 0])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _missing_header_workbook() -> bytes:
    workbook = openpyxl.Workbook()
    monthly = workbook.active
    monthly.title = "Janeiro"
    monthly.append(["not a recognised header"])
    summary = workbook.create_sheet("Balanço final")
    summary.append([2026])
    summary.append(["Mês", "Person Alpha"])
    for _ in range(12):
        summary.append(["Mês", 0])
    summary.append(["Total", 0])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _malformed_xlsx_bytes() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"<Types>")
        archive.writestr("xl/workbook.xml", b"<workbook>")
    return buffer.getvalue()


def _admin_context(db, email: str = ADMIN) -> AuthContext:
    user = db.query(User).filter(User.email == email).one()
    return AuthContext(
        user=user,
        person_id=user.person_id,
        role_codes=frozenset({"administrador", "chefe_operacoes"}),
        permission_codes=frozenset({"time_bank.view_own", "time_bank.view_all", "time_bank.manage_all"}),
    )


def _pending_batch(db, ctx: AuthContext) -> TimeBankImportBatch:
    batch = TimeBankImportBatch(
        sha256=uuid.uuid4().hex * 2,
        parser_version=PARSER_VERSION,
        status="pending",
        source_year=2026,
        parsed_payload={
            "entries": [],
            "openings": [],
            "people": [],
            "issues": [],
            "counts": {"entries": 0, "openings": 0, "needs_review": 0},
        },
        resolutions={},
        created_by_user_id=ctx.user.id,
        created_at=dt.datetime.now(dt.timezone.utc),
    )
    db.add(batch)
    db.flush()
    return batch


def _post_preview_without_server_raise(api_client, content: bytes, filename: str = "upload.xlsx"):
    # The fixture keeps the real dependency override/session but normally raises
    # application exceptions in the test process. HTTP validation must be tested
    # as an actual response, not as a Python exception escaping TestClient.
    with TestClient(app, raise_server_exceptions=False) as client:
        return client.post(
            "/api/time-bank/imports/preview",
            headers=UPLOAD_HEADERS,
            files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )


@pytest.mark.parametrize(
    ("filename", "content", "expected_detail"),
    [
        ("not-xlsx.bin", b"not-an-xlsx", "O ficheiro não é um XLSX válido."),
        ("oversized.xlsx", b"x" * (MAX_UPLOAD_BYTES + 1), "O ficheiro excede o limite de 5 MiB."),
        ("malformed.xlsx", _malformed_xlsx_bytes(), "Não foi possível ler a estrutura XLSX."),
    ],
)
def test_invalid_uploads_are_sanitized_422_and_create_no_staging_row(
    db_session, api_client, filename: str, content: bytes, expected_detail: str
):
    before = db_session.query(TimeBankImportBatch).count()

    response = _post_preview_without_server_raise(api_client, content, filename)

    assert response.status_code == 422, response.text
    assert response.json() == {"detail": expected_detail}
    assert "Internal Server Error" not in response.text
    assert db_session.query(TimeBankImportBatch).count() == before


def test_preview_http_matrix_keeps_structural_issues_pending_and_does_not_evaluate_unsafe_formula(
    db_session, api_client
):
    missing_header = _post_preview_without_server_raise(api_client, _missing_header_workbook(), "missing-header.xlsx")
    assert missing_header.status_code == 200, missing_header.text
    assert any(issue["code"] == "missing_daily_header" and issue["blocking"] for issue in missing_header.json()["issues"])

    unsafe = _post_preview_without_server_raise(
        api_client, _workbook_bytes(unsafe_formula=True), "unsafe-formula.xlsx"
    )
    assert unsafe.status_code == 200, unsafe.text
    unsafe_body = unsafe.json()
    assert any(issue["code"] == "unknown_formula" and issue["blocking"] for issue in unsafe_body["issues"])
    assert "WEBSERVICE" not in json.dumps(unsafe_body, ensure_ascii=False)

    valid = _post_preview_without_server_raise(api_client, _workbook_bytes(), "valid.xlsx")
    assert valid.status_code == 200, valid.text
    assert valid.json()["status"] == "pending"
    assert valid.json()["counts"]["entries"] == 1


def test_preview_rejects_new_caller_state_before_parsing_without_flushing_or_rolling_back(db_session):
    ctx = _admin_context(db_session)
    before = db_session.query(TimeBankImportBatch).count()
    unrelated = Person(display_name="Unrelated synthetic pending", email=None, is_active=True)
    db_session.add(unrelated)

    with pytest.raises(TimeBankError, match="sessão.*limpa"):
        preview_import(db_session, ctx, b"not parsed because session is dirty")

    assert unrelated in db_session.new
    assert db_session.query(TimeBankImportBatch).count() == before
    assert unrelated.display_name == "Unrelated synthetic pending"


def test_resolve_rejects_dirty_caller_state_before_lock_without_flushing_or_rollback(db_session):
    ctx = _admin_context(db_session)
    batch = _pending_batch(db_session, ctx)
    person = db_session.query(Person).filter(Person.display_name == "PM Sintético Um").one()
    original_name = person.display_name
    person.display_name = "Changed only inside caller session"

    with pytest.raises(TimeBankError, match="sessão.*limpa"):
        resolve_import(db_session, ctx, batch.id, {})

    assert person in db_session.dirty
    assert person.display_name == "Changed only inside caller session"
    assert db_session.get(TimeBankImportBatch, batch.id).resolutions == {}
    person.display_name = original_name


def test_preview_can_be_resolved_in_same_session_without_rejecting_its_own_audit_row(db_session):
    ctx = _admin_context(db_session)
    label = f"Chained-{uuid.uuid4().hex}"
    batch = preview_import(db_session, ctx, _workbook_bytes(label=label))
    person = db_session.query(Person).filter(Person.display_name == "PM Sintético Um").one()

    resolved = resolve_import(db_session, ctx, batch.id, {label: person.id})

    assert resolved.resolutions == {label: str(person.id)}


def test_apply_rejects_deleted_caller_state_before_lock_without_flushing_or_rollback(db_session):
    ctx = _admin_context(db_session)
    batch = _pending_batch(db_session, ctx)
    unrelated = Person(display_name="Unrelated synthetic deletable", email=None, is_active=True)
    db_session.add(unrelated)
    db_session.flush()
    db_session.delete(unrelated)

    with pytest.raises(TimeBankError, match="sessão.*limpa"):
        apply_import(db_session, ctx, batch.id)

    assert unrelated in db_session.deleted
    assert db_session.get(TimeBankImportBatch, batch.id).status == "pending"


@pytest.mark.skipif(engine.dialect.name != "postgresql", reason="requires a real PostgreSQL engine")
def test_postgresql_concurrent_same_hash_preview_returns_one_pending_batch_to_both_manager_sessions(db_session):
    content = _workbook_bytes(label=f"Concurrent-{uuid.uuid4().hex}")
    before_entries = db_session.query(TimeBankEntry).count()
    before_adjustments = db_session.query(TimeBankAdjustment).count()
    before_batches = db_session.query(TimeBankImportBatch).count()

    barrier = threading.Barrier(2, timeout=15)
    armed: set[int] = set()
    armed_lock = threading.Lock()
    outcomes: list[tuple[str, str, str]] = []
    outcome_lock = threading.Lock()

    def before_flush(session, _flush_context, _instances):
        key = id(session)
        with armed_lock:
            wait = key in armed and any(isinstance(item, TimeBankImportBatch) for item in session.new)
            if wait:
                armed.remove(key)
        if wait:
            barrier.wait()

    event.listen(OrmSession, "before_flush", before_flush)

    def worker(email: str) -> None:
        db = SessionLocal()
        try:
            ctx = _admin_context(db, email)
            with armed_lock:
                armed.add(id(db))
            batch = preview_import(db, ctx, content)
            db.commit()
            with outcome_lock:
                outcomes.append(("ok", batch.status, str(batch.id)))
        except BaseException as exc:  # assertion below reports any unexpected exception
            with outcome_lock:
                outcomes.append(("error", type(exc).__name__, str(exc)))
            db.rollback()
        finally:
            db.close()

    threads = [threading.Thread(target=worker, args=(ADMIN,)), threading.Thread(target=worker, args=(CHEFE,))]
    try:
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=30)
    finally:
        event.remove(OrmSession, "before_flush", before_flush)

    assert not any(thread.is_alive() for thread in threads), outcomes
    assert len(outcomes) == 2, outcomes
    assert all(status == "ok" for status, _value, _detail in outcomes), outcomes
    assert {value for _status, value, _detail in outcomes} == {"pending"}, outcomes
    assert len({detail for _status, _value, detail in outcomes}) == 1, outcomes
    assert db_session.query(TimeBankImportBatch).count() == before_batches + 1
    assert db_session.query(TimeBankEntry).count() == before_entries
    assert db_session.query(TimeBankAdjustment).count() == before_adjustments

    # The two worker sessions commit outside this test's external transaction.
    # Remove their synthetic batch so later tests remain isolated regardless of
    # pytest's module ordering (the SQLite variant is skipped above).
    worker_batch_ids = [
        uuid.UUID(detail)
        for status, value, detail in outcomes
        if status == "ok" and value == "pending"
    ]
    cleanup_session = SessionLocal()
    try:
        cleanup_session.query(TimeBankImportBatch).filter(TimeBankImportBatch.id.in_(worker_batch_ids)).delete(
            synchronize_session=False
        )
        cleanup_session.commit()
    finally:
        cleanup_session.close()
