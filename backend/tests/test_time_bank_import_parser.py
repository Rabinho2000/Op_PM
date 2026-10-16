from __future__ import annotations

import datetime as dt
import io

import openpyxl
from openpyxl.comments import Comment

from app.services.time_bank_import import parse_time_bank_xlsx


def _workbook_bytes(*, unknown_formula: bool = False) -> bytes:
    workbook = openpyxl.Workbook()
    january = workbook.active
    january.title = "Janeiro"
    january.append(["Janeiro"])
    january.append(["Colaborador", "Weekday", "Data", "Hora de entrada", "Hora de almoço", "Hora de saída", "Balanço"])
    january.append(["Person Alpha", 2, dt.datetime(2026, 1, 5), 9, "sim", 18.5, 0])
    january.append(["Person Beta", 6, dt.datetime(2026, 1, 9), None, None, None, 1.5])
    if unknown_formula:
        january["G3"] = "=WEBSERVICE(\"https://example.invalid\")"
    summary = workbook.create_sheet("Balanço final")
    summary.append([2026])
    summary.append(["Mês", "Person Alpha", "Person Beta"])
    for month in range(12):
        summary.append(["Mês", 0, 0])
    summary.append(["Total", 0, 0])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_parser_extracts_synthetic_daily_rows_without_evaluating_formulas():
    result = parse_time_bank_xlsx(_workbook_bytes())
    assert result["source_year"] == 2026
    assert result["counts"]["entries"] == 2
    assert result["counts"]["openings"] == 0
    assert result["source_hash"]
    assert {row["source_label"] for row in result["entries"]} == {"Person Alpha", "Person Beta"}
    assert result["entries"][1]["needs_review"] is True
    assert not any(issue["blocking"] for issue in result["issues"])


def test_unknown_formula_without_cache_is_blocking_and_never_becomes_zero():
    result = parse_time_bank_xlsx(_workbook_bytes(unknown_formula=True))
    assert any(issue["code"] == "unknown_formula" and issue["blocking"] for issue in result["issues"])
    formula_row = result["entries"][0]
    assert formula_row["balance_seconds"] is None


def test_parser_keeps_labeled_undated_zero_row_in_quarantine():
    workbook = openpyxl.Workbook()
    january = workbook.active
    january.title = "Janeiro"
    january.append(["Janeiro"])
    january.append(["Colaborador", "Weekday", "Data", "Hora de entrada", "Hora de almoço", "Hora de saída", "Balanço"])
    january.append(["Person Alpha", 2, dt.datetime(2026, 1, 5), 9, "sim", 18.5, 0])
    january.append(["Person Alpha", 0, None, None, None, None, 0])
    summary = workbook.create_sheet("Balanço final")
    summary.append([2026])
    summary.append(["Mês", "Person Alpha"])
    for _ in range(12):
        summary.append(["Mês", 0])
    summary.append(["Total", 0])
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = parse_time_bank_xlsx(buffer.getvalue())

    assert result["counts"]["source_rows"] == 2
    assert result["counts"]["entries"] == 1
    assert result["counts"]["quarantined"] == 1
    assert result["quarantined"][0]["source_label"] == "Person Alpha"
    assert result["quarantined"][0]["balance_seconds"] == 0
    assert result["quarantined"][0]["work_date"] is None
    assert any(issue["code"] == "undated_zero_quarantined" and not issue["blocking"] for issue in result["issues"])


def test_parser_blocks_undated_nonzero_row_without_fabricating_a_date():
    workbook = openpyxl.Workbook()
    january = workbook.active
    january.title = "Janeiro"
    january.append(["Janeiro"])
    january.append(["Colaborador", "Weekday", "Data", "Hora de entrada", "Hora de almoço", "Hora de saída", "Balanço"])
    january.append(["Person Alpha", 0, None, None, None, None, 1.5])
    summary = workbook.create_sheet("Balanço final")
    summary.append([2026])
    summary.append(["Mês", "Person Alpha"])
    for _ in range(12):
        summary.append(["Mês", 0])
    summary.append(["Total", 0])
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = parse_time_bank_xlsx(buffer.getvalue())

    assert result["counts"]["entries"] == 0
    assert result["counts"]["quarantined"] == 1
    assert result["quarantined"][0]["balance_seconds"] == 5400
    assert result["quarantined"][0]["work_date"] is None
    assert any(issue["code"] == "undated_nonzero" and issue["blocking"] for issue in result["issues"])


def test_parser_preserves_safe_comments_and_redacts_secret_shaped_notes():
    workbook = openpyxl.Workbook()
    january = workbook.active
    january.title = "Janeiro"
    january.append(["Janeiro"])
    january.append(["Colaborador", "Weekday", "Data", "Hora de entrada", "Hora de almoço", "Hora de saída", "Balanço", "Nota"])
    january.append(["Person Alpha", 2, dt.datetime(2026, 1, 5), 9, "sim", 18.5, 0, "Nota fora da tabela"])
    january["G3"].comment = Comment(
        "Comentário preservado; https://example.invalid/token?secret=redact", "synthetic"
    )
    summary = workbook.create_sheet("Balanço final")
    summary.append([2026])
    summary.append(["Mês", "Person Alpha"])
    for _ in range(12):
        summary.append(["Mês", 0])
    summary.append(["Total", 0])
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = parse_time_bank_xlsx(buffer.getvalue())
    row = result["entries"][0]

    assert "Nota fora da tabela" in row["reason"]
    assert "https://example.invalid" not in row["reason"]
    assert row["review_flags"]["note_redacted"] is True
    assert any(issue["code"] == "note_redacted" and not issue["blocking"] for issue in result["issues"])


class _Cell:
    def __init__(self, value):
        self.value = value


class _SummarySheet:
    title = "Balanço final"

    def __init__(self, row1, row2, row15, formulas, totals):
        self._rows = {
            1: [_Cell(value) for value in row1],
            2: [_Cell(value) for value in row2],
            15: [_Cell(value) for value in row15],
        }
        self._formulas = [_Cell(value) for value in formulas]
        self._totals = totals

    def __getitem__(self, key):
        if isinstance(key, int):
            if key == 15 and self._formulas:
                return self._formulas
            return self._rows[key]
        column = key.removesuffix("15")
        return _Cell(self._totals[column])


def test_openings_only_use_positive_references_into_approved_previous_year_block():
    from app.services.time_bank_import import _parse_openings

    row1 = [2026, None, None, None, None, None, None, None, None, 2025, None, None, None, None, None, None, 2024]
    row2 = ["Mês", "Current A", "Current B", "Current C", None, None, None, None, None, "Total", "Old A", "Old B", None, None, None, None, "Total"]
    cached_row15 = ["Total", 0, 0, 0, None, None, None, None, None, "Total", 2, 0, None, None, None, None, "Total"]
    formula_row15 = [
        "Total",
        "=SUM(B3:B14)+K15",
        "=SUM(C3:C14)+Q15",
        "=SUM(D3:D14)-K15",
        None,
        None,
        None,
        None,
        None,
        "Total",
        "2",
        "0",
        None,
        None,
        None,
        None,
        "Total",
    ]
    formula = _SummarySheet(row1, row2, cached_row15, formula_row15, {"K": 2, "Q": 10})
    cached = _SummarySheet(row1, row2, cached_row15, [], {"K": 2, "Q": 10})
    issues = []

    openings = _parse_openings(formula, cached, file_hash="hash", source_year=2026, issues=issues)

    assert [(row["source_label"], row["delta_seconds"]) for row in openings] == [("Current A", 7200)]
    assert not any(issue.blocking for issue in issues)
